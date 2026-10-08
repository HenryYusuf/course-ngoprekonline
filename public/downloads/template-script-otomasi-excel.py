#!/usr/bin/env python3
"""Template otomasi Excel, paket artikel "Excel Unleashed" (Ngoprek.Online).

Cara pakai:
  1. pip install openpyxl
  2. python template-script-otomasi-excel.py data.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def ringkas(path: Path) -> None:
    """Baca workbook, cetak peta sheet dan 5 baris pertama tiap sheet."""
    if not path.exists():
        raise SystemExit(f"File tidak ditemukan: {path}")

    wb = load_workbook(path, data_only=True)
    print(f"Workbook : {path.name}")
    print(f"Sheets   : {', '.join(wb.sheetnames)}")

    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n[{name}] {ws.max_row} baris x {ws.max_column} kolom")
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            print("  ", row)


def tambah_kolom_total(path: Path, keluar: Path) -> None:
    """Contoh tulis: tambahkan kolom total di lembar pertama, simpan ke file baru."""
    wb = load_workbook(path)
    ws = wb.active
    ws.cell(row=1, column=ws.max_column + 1, value="Total")
    total_col = ws.max_column

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=total_col, value=f"=B{row}*C{row}")

    wb.save(keluar)
    print(f"Hasil disimpan ke: {keluar}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Pakai: python template-script-otomasi-excel.py data.xlsx")

    src = Path(sys.argv[1])
    ringkas(src)
    tambah_kolom_total(src, src.with_name(f"{src.stem}-otomatis{src.suffix}"))


if __name__ == "__main__":
    main()
